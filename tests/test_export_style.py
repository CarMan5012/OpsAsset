from types import SimpleNamespace

from openpyxl import load_workbook

from app.services.excel_service import export_hosts_data


def make_host(name: str, ip: str, notes: str = ""):
    return SimpleNamespace(
        hostname=name,
        private_ip=ip,
        public_ip="",
        open_ports="6379",
        env="prod",
        status="online",
        cpu_cores=4,
        memory_gb=32,
        disk_gb=100,
        arch="amd64",
        os="CentOS 7.7",
        kernel_version="3.10.0",
        cluster_relations=[],
        notes=notes,
        created_at=None,
        updated_at=None,
    )


def test_export_matches_centered_report_layout():
    hosts = [make_host("node-01", "10.0.0.1"), make_host("node-02", "10.0.0.2", "长备注" * 40)]
    stream, _, _ = export_hosts_data(
        hosts,
        selected_columns=["hostname", "private_ip", "public_ip", "open_ports", "env", "status", "clusters", "notes"],
    )
    ws = load_workbook(stream).active

    assert ws.column_dimensions["C"].width == 21
    assert ws.column_dimensions["G"].width == 19
    assert ws.column_dimensions["H"].width == 29
    assert ws.row_dimensions[1].height == 27
    assert ws.row_dimensions[2].height == 34
    assert ws.row_dimensions[3].height > 34
    assert ws["A1"].fill.fgColor.rgb.endswith("E9EEF3")
    assert ws["A1"].font.name == "DengXian"
    assert ws["A1"].font.sz == 10.5
    assert ws["A1"].font.bold is True
    assert ws["A1"].alignment.horizontal == "center"
    assert ws["B2"].font.name == "DengXian"
    assert ws["B2"].font.sz == 10
    assert ws["B2"].alignment.horizontal == "center"
    assert ws["B2"].border.left.style == "thin"
    assert ws["B2"].border.right.style == "thin"
    assert ws["B2"].font.bold is not True
    assert ws["D2"].number_format == "@"
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:H3"
    hidden_filter_columns = {column.colId for column in ws.auto_filter.filterColumn if column.hiddenButton}
    assert hidden_filter_columns == {0, 1, 2, 3, 7}
    assert all(not column.showButton for column in ws.auto_filter.filterColumn)


def test_export_filter_buttons_follow_selected_column_order():
    stream, _, _ = export_hosts_data(
        [make_host("node-01", "10.0.0.1")],
        selected_columns=["notes", "status", "hostname", "os", "cpu_cores", "env", "arch", "clusters", "memory_gb", "disk_gb"],
    )
    ws = load_workbook(stream).active

    assert ws.auto_filter.ref == "A1:J2"
    visible_headers = {
        ws.cell(1, col_idx + 1).value
        for col_idx in range(ws.max_column)
        if col_idx not in {column.colId for column in ws.auto_filter.filterColumn if column.hiddenButton}
    }
    assert visible_headers == {"环境", "状态", "CPU", "内存", "数据盘", "操作系统", "所属服务"}
